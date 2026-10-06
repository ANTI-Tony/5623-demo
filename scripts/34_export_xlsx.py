r"""Export every experimental result and the full configuration to one workbook.

Three kinds of sheet:
  CONFIG        every setting that determines a number, with the file it lives in
  <table>       one per list-shaped result file
  summary_stats every scalar the paper reports, flattened, with its source

Configuration is read from the code where it can be, and hand-transcribed with a
source reference where it cannot (a constant inside a function, a fact recorded in
the runbook). Anything transcribed names its file so it can be checked.

  out: UndoAtlas_data.xlsx
"""
from __future__ import annotations
import json, re, collections
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]; OUT = ROOT / 'out'
read = lambda p: (ROOT / p).read_text() if (ROOT / p).exists() else ''

# ---------------------------------------------------------------- configuration
def grid_config():
    src = read('scripts/30_e1_grid.py')
    verbs = re.findall(r"^\s*'(\w+)':\s*dict\(seq=", src, re.M)
    states = re.findall(r"^\s*'(\w+)':\s*dict\(recipe=", src, re.M)
    lex = re.search(r"INVERSE_LEXICON = \((.*?)\)", src, re.S)
    return verbs, states, re.findall(r"'([^']+)'", lex.group(1)) if lex else []

VERBS, STATES, LEXICON = grid_config()
TASKS = sorted({r['task'] for r in json.loads((OUT / 'k2.json').read_text()) if 'task' in r}) \
    if (OUT / 'k2.json').exists() else []
PIN = {k: v for k, v in (l.split('==') for l in read('requirements-py312.txt').splitlines()
                         if '==' in l and l.split('==')[0] in
                         ('browsergym-core', 'browsergym-webarena-verified', 'playwright',
                          'webarena-verified', 'anthropic', 'openpyxl'))}

CONFIG = [
    ('model', 'model id', 'claude-sonnet-5', 'cawebagent/lm/protocol.py:51 (env CAWE_MODEL)'),
    ('model', 'second model, cross-model arm', 'claude-sonnet-4-6', 'main.tex, cross-model appendix'),
    ('model', 'price in / out, USD per Mtok', '2.00 / 10.00', 'cawebagent/lm/protocol.py:21'),
    ('model', 'max_tokens per call', '1500', 'cawebagent/lm/protocol.py (output cap; 14 steps hit it)'),
    ('effort', 'T0 — low arm', 'effort=low, thinking=disabled', 'cawebagent/lm/protocol.py TIERS'),
    ('effort', 'T3 — high arm', 'effort=xhigh, thinking=adaptive', 'cawebagent/lm/protocol.py TIERS'),
    ('effort', 'tiers defined but unused in the paper', 'T1 medium, T2 high, T3max max', 'cawebagent/lm/protocol.py TIERS'),
    ('episode', 'step cap, primary runs', '12', 'scripts/05_batch.py --steps default'),
    ('episode', 'step cap, budget sweep', '20', 'scripts/15_step_budget.py'),
    ('episode', 'arms per task', 'T0, T3', 'scripts/05_batch.py --tiers default'),
    ('episode', 'seeds', '2 (seed 1, seed 2 identical protocol)', 'runs/batch, runs/batch_seed2'),
    ('episode', 'tasks', f'{len(TASKS)} tier-2 shopping_admin tasks: {TASKS}', 'out/k2.json'),
    ('episode', 'spend cap enforced in-tool', 'USD 12.00 default per batch', 'scripts/05_batch.py --cap'),
    ('environment', 'application', 'Magento admin (WebArena shopping_admin), self-hosted container', 'cawebagent/trace/reset.py'),
    ('environment', 'reset method', 'ContainerManager.start() — destroys and recreates; docker restart does NOT roll back the DB', 'cawebagent/trace/reset.py'),
    ('environment', 'reset time', '~16 s warm, ~175 s cold', 'measured; RUNBOOK.md'),
    ('environment', 'port', '7780 (admin at /admin)', 'scripts/30_e1_grid.py ADMIN'),
    ('capture', 'HAR content mode', 'embed — required, 447 evaluators read post bodies', 'cawebagent/trace/runner.py:53'),
    ('capture', 'HAR mode', 'full', 'cawebagent/trace/runner.py:54'),
    ('capture', 'browser', 'Chromium via Playwright, headless', 'cawebagent/trace/runner.py:41'),
    ('capture', 'per-action timeout', '30000 ms', 'cawebagent/trace/runner.py:41'),
    ('capture', 'dialog handling', 'auto-accept; Playwright dismisses by default, which silently cancels the POST', 'cawebagent/trace/runner.py'),
    ('capture', 'phase segmentation', 'setup / agent / validate, by runner wall-clock marks', 'cawebagent/trace/har_parse.py'),
    ('observation', 'accessibility tree pruning cap', '14000 chars', 'cawebagent/agent/prune.py:14'),
    ('observation', 'static text retained near actionable', '1 line either side', 'cawebagent/agent/prune.py:14'),
    ('grid', 'write verbs', ', '.join(VERBS), 'scripts/30_e1_grid.py VERBS'),
    ('grid', 'object states', ', '.join(STATES), 'scripts/30_e1_grid.py STATES'),
    ('grid', 'cells', f'{len(VERBS)} verbs x {len(STATES)} states = {len(VERBS)*len(STATES)}', 'scripts/30_e1_grid.py'),
    ('grid', 'positive control, re-run every batch', 'hold x pending, inverse = Unhold', 'scripts/30_e1_grid.py CONTROL_CELL'),
    ('grid', 'inverse-matching lexicon', ', '.join(LEXICON), 'scripts/30_e1_grid.py INVERSE_LEXICON'),
    ('grid', 'cells per container reset', '10 default, 1 under --isolate', 'scripts/30_e1_grid.py --batch'),
    ('grid', 'outcome decided by', 'state predicate read back from the page — never the click return, which always reports a timeout and still lands', 'scripts/30_e1_grid.py NFR8 rationale'),
    ('statistics', 'permutation draws', '10000', 'scripts/31_abstract_headline.py DRAWS'),
    ('statistics', 'permutation RNG seed', '0', 'scripts/31_abstract_headline.py RNG'),
    ('statistics', 'permutation null (reported)', 'NULL-A: shuffle commit labels within episode, recompute the neighbourhood', 'scripts/31_abstract_headline.py'),
    ('statistics', 'permutation null (named, not reported)', 'NULL-B: shuffle neighbourhood membership directly', 'scripts/31_abstract_headline.py'),
    ('statistics', 'bootstrap draws, kappa CI', '20000', 'scripts/33_calibration_and_margin.py DRAWS'),
    ('statistics', 'bootstrap RNG seed', '0', 'scripts/33_calibration_and_margin.py RNG'),
    ('statistics', 'paired outcome test', 'exact McNemar, two-sided', 'scripts/09_stats.py'),
    ('statistics', 'proportion intervals', 'Wilson', 'scripts/09_stats.py'),
    ('statistics', '2x2 independence', "Fisher exact", 'scripts/09_stats.py'),
    ('exclusions', 'HTTP archives', 'excluded from the artifact: 3.3 GB against 3 MB for everything else', '.gitignore, README.md'),
    ('exclusions', 'GitLab site', 'not hosted: 22 GB compressed image; carries 43 of 89 tier-2 events', 'main.tex limitations'),
]
for k, v in sorted(PIN.items()):
    CONFIG.append(('versions', k, v, 'requirements-py312.txt'))
CONFIG.append(('versions', 'python', '3.12', 'pyproject.toml requires-python'))

# ---------------------------------------------------------------- result tables
TABLES = [
    ('reversibility_grid', 'e1_grid.json', 'Every (verb x object state) cell: outcome, inverse taken, the warning rendered at the gate frame, rejected candidates', '30_e1_grid.py'),
    ('steps_pooled', 'k3_steps_pooled.json', 'Per agent step over two seeds: commit indicator, output tokens, input tokens, cost', '08_instrument.py'),
    ('steps_seed1', 'k3_steps.json', 'Per step, seed 1 only', '08_instrument.py'),
    ('steps_seed2', 'k3_steps_batch_seed2.json', 'Per step, seed 2 only', '08_instrument.py'),
    ('steps_model2', 'k3_steps_batch_sonnet46.json', 'Per step, cross-model arm', '08_instrument.py'),
    ('taxonomy', 'exp1_taxonomy.json', 'All benchmark tasks with declared state-changing events and consequence tier', 'exp1_consequence_taxonomy.py'),
    ('declared_events', 'exp0_events.json', 'Every expected network event the benchmark declares', 'exp0_dataset_audit.py'),
    ('site_structure', 'k13_site_structure.json', 'Per site: distinct state-changing endpoints and the share carrying an entity id', '26_site_structure.py'),
    ('gate_requests', 'gate_check.json', 'Every recorded state-changing request with each pre-commit check verdict', '20_precommit_check.py'),
    ('gate_dataset', 'gate_dataset.json', 'The request corpus the gate table is computed over', '19_build_gate_dataset.py'),
    ('upstream_attribution', 'k2_labels.json', 'Committed-but-wrong episodes and where the deciding error was made', '22_k2_labels.py'),
    ('upstream_sharp', 'k2_sharp.json', 'The same episodes under the sharpened attribution rule', '07b_k2_sharp.py'),
    ('episodes', 'k2.json', 'Per episode: task, arm, score, collateral count, steps, cost', '07_k2.py'),
    ('step_events', 'k3_events.json', 'Classified state-changing events joined to the step window that emitted them', '08c_events.py'),
    ('undo_probe', 'undo_report.json', 'The original four-class irreversibility probe with its positive control', '11_undo_probe.py'),
    ('control_tier01', 'control_tier01_detail.json', 'The tier-0/1 control arm selection', '13_select_control.py'),
    ('reattribution', '../c3_reattribution.json', 'Commits re-attributed from the confirmation click to the decision step', '24_c3_reattribution.py'),
    ('order_pool', 'e1_order_pool.json', 'Order id to baseline status at the image baseline', '30_e1_grid.py'),
    ('gold_sheet', 'gold_sheet.csv', 'The hand-annotation sheet for recall against an independent standard (annotation column not yet filled)', '28_build_gold_sheet.py'),
]
SCALARS = [
    ('headline_stats', 'k3_headline.json', 'Compute-to-change statistics under both permutation nulls', '31_abstract_headline.py'),
    ('calibration', 'k16_calibration.json', 'Warning-calibration kappa with bootstrap CI, and the phase-boundary margin', '33_calibration_and_margin.py'),
    ('two_seed', 'k7_two_seed.json', 'Effort-arm comparison at both seeds and the run-to-run flip rate', '14_two_seed.py'),
    ('step_budget', 'k8_step_budget.json', 'What raising the step cap from 12 to 20 rescues', '15_step_budget.py'),
    ('r2_controls', 'k9_r2_controls.json', 'Controls on the compute measurement: output-cap contamination, task clustering, predecessor matching', '16_r2_controls.py'),
    ('oracle_cost', 'k10_oracle.json', 'What a consequence-conditioned spend would have cost, as an oracle bound', '17_oracle_allocator.py'),
    ('cross_model', 'k11_cross_model.json', 'The cross-model arm and its parse-failure rate', '18_cross_model.py'),
    ('detector_escapes', 'k12_recall.json', 'Routes by which a state change escapes the detector, counted over all episodes', '21_detector_recall.py'),
    ('state_undo', 'k14_state_undo.json', 'One order driven through three states, recording what the interface offers at each', '27_state_dependent_undo.py'),
    ('blindness', 'exp0b_blindness.json', 'How many benchmark evaluators can assert that something should NOT happen', 'exp0b_blindness.py'),
    ('dataset_audit', 'exp0_tasks.json', 'Benchmark composition: task counts by type, evaluator counts', 'exp0_dataset_audit.py'),
    ('cross_benchmark', 'k6_cross_benchmark.json', 'The same audit against the original WebArena schema', '12_cross_benchmark.py'),
    ('baserate', 'k2_baserate.json', 'Base rates for the positional statistic', '07c_k2_baserate.py'),
    ('gate_summary', 'gate_check.json', 'Interruption, recall and precision per check', '23_id_rule.py'),
    ('t4_demo', 't4_result.json', 'The two-arm demonstration: identical official score, different damage', 'm2_t4_collateral.py'),
    ('sanity_cell', 'sanity_new_cell.json', 'The first new grid cell, with end-to-end timings', 'sanity_new_cell.py'),
]

HDR = Font(bold=True, color='FFFFFF'); FILL = PatternFill('solid', fgColor='4A5568')
def style(ws, ncol, widths=None):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c); cell.font, cell.fill = HDR, FILL
        cell.alignment = Alignment(vertical='top')
        w = (widths or {}).get(c) or min(60, max(12, max(
            (len(str(ws.cell(row=r, column=c).value or '')) for r in range(1, min(ws.max_row, 80) + 1)), default=12) + 2))
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = 'A2'
flat = lambda v: json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v

wb = Workbook(); readme = wb.active; readme.title = 'README'

ws = wb.create_sheet('CONFIG')
ws.append(['group', 'setting', 'value', 'source'])
for row in CONFIG: ws.append(list(row))
style(ws, 4, {1: 14, 2: 34, 3: 62, 4: 46})

made = [('CONFIG', len(CONFIG), 'Every setting that determines a number, with the file it lives in', '—')]
def load_rows(f):
    """Return (rows, cols) for anything we ship: a CSV, a list of records, a bare
    list, or a map. Returns (None, None) when there is nothing to tabulate."""
    if f.suffix == '.csv':
        lines = [l for l in f.read_text().splitlines() if l.strip()]
        if len(lines) < 2: return None, None
        cols = lines[0].split(',')
        return [dict(zip(cols, l.split(','))) for l in lines[1:]], cols
    d = json.loads(f.read_text())
    if not d: return None, None
    if isinstance(d, dict):
        return [{'key': k, 'value': flat(v)} for k, v in d.items()], ['key', 'value']
    if not isinstance(d[0], dict):
        return [{'value': flat(x)} for x in d], ['value']
    return d, list(dict.fromkeys(k for r in d for k in r))

for name, fn, what, script in TABLES:
    f = (OUT / fn).resolve()
    if not f.exists():
        made.append((name, 0, what + '  [MISSING]', script)); continue
    rows, cols = load_rows(f)
    if rows is None:
        made.append((name, 0, what + '  [EMPTY]', script)); continue
    sh = wb.create_sheet(name[:31]); sh.append(cols)
    for r in rows: sh.append([flat(r.get(c)) for c in cols])
    style(sh, len(cols)); made.append((name, len(rows), what, script))

ws = wb.create_sheet('summary_stats')
ws.append(['group', 'key', 'value', 'what it is', 'produced by'])
seen = set()
for name, fn, what, script in SCALARS:
    f = OUT / fn
    if not f.exists() or fn in seen: continue
    d = json.loads(f.read_text())
    def walk(prefix, obj):
        if isinstance(obj, dict):
            for k, v in obj.items(): walk(f'{prefix}.{k}' if prefix else k, v)
        elif isinstance(obj, list) and obj and not isinstance(obj[0], dict):
            ws.append([name, prefix, flat(obj), what, script])
        elif not isinstance(obj, list):
            ws.append([name, prefix, obj, what, script])
    walk('', d)
style(ws, 5, {3: 46, 4: 54, 5: 30}); made.append(('summary_stats', ws.max_row - 1, 'Every scalar the paper reports, flattened, with its source', 'various'))

readme.append(['sheet', 'rows', 'what it is', 'produced by'])
for row in made: readme.append(list(row))
style(readme, 4, {1: 22, 2: 8, 3: 78, 4: 30})
readme.insert_rows(1, 2)
readme['A1'] = ('UndoAtlas — all experimental results and the full configuration. '
                'CONFIG holds every setting that determines a number. HTTP archives are excluded (3.3 GB); '
                'every sheet regenerates from out/ via scripts/34_export_xlsx.py.')
readme['A1'].font = Font(bold=True, size=12)

p = ROOT / 'UndoAtlas_data.xlsx'; wb.save(p)
print(f'wrote {p}\n')
for name, rows, *_ in made: print(f'  {name:<24}{rows:>6}')
